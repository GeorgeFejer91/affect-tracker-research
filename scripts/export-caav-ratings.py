"""Export the complete official CAAV rating workbooks to stable UTF-8 CSV."""

from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RATINGS = ROOT / "site" / "assets" / "research-stimuli" / "caav" / "v1" / "ratings"
ORIGINAL = RATINGS / "original"


def export_rows(path: Path, rows: list[tuple[object, ...]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerows(rows)


def nonempty_rows(worksheet) -> list[tuple[object, ...]]:
    return [row for row in worksheet.iter_rows(values_only=True) if any(value is not None for value in row)]


def main() -> None:
    RATINGS.mkdir(parents=True, exist_ok=True)

    norms = load_workbook(ORIGINAL / "CAAV_Dataset.xlsx", read_only=True, data_only=True)["CAAV_Rating"]
    norm_rows = nonempty_rows(norms)
    if norm_rows[0][0] != "CAAV Dataset" or norm_rows[1][0] != "Video_ID":
        raise ValueError("Unexpected CAAV aggregate workbook structure")
    aggregate = [norm_rows[1], *[row for row in norm_rows[2:] if re.fullmatch(r"[13]_[FM]_\d{3}", str(row[0]))]]
    if len(aggregate) != 361:
        raise ValueError(f"Expected one header and 360 aggregate rows, got {len(aggregate)}")
    export_rows(RATINGS / "caav__clip-norms__sam-1-9.csv", aggregate)

    raw = load_workbook(ORIGINAL / "CAAV_RawData.xlsx", read_only=True, data_only=True)
    for sheet_name, dimension in (("Valence", "valence"), ("Arousal", "arousal")):
        rows = nonempty_rows(raw[sheet_name])
        if rows[0][:4] != ("Sub", "Age", "Gender", "Sub_List"):
            raise ValueError(f"Unexpected CAAV {sheet_name} workbook structure")
        export_rows(RATINGS / f"caav__raw-{dimension}__sam-1-9.csv", rows)

    print("Exported complete CAAV aggregate, raw-valence, and raw-arousal ratings")


if __name__ == "__main__":
    main()
